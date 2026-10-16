"""Verify the downloaded workbook, including Excel cell types and financial ties."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.routers.valuation import sample_data

client = TestClient(app)


def inputs(years=5):
    return dict(historical_base=sample_data(), projection_years=years,
        operating_assumptions={'revenue_growth_rates': [.1]*years, 'cogs_percent_revenue': [.4]*years},
        wc_drivers={'days_sales_outstanding': 45, 'days_inventory_outstanding': 60, 'days_payable_outstanding': 40},
        capex_assumptions={'growth_capex': [.5]+[0]*(years-1)}, financing_assumptions={})


def row(sheet, label):
    return next(c.row for c in sheet['B'] if c.value == label)


@pytest.mark.parametrize('years', [1, 5, 30])
def test_download_is_numeric_readable_and_reconciles(years):
    response = client.post('/api/export/excel/three-statement', json=inputs(years))
    assert response.status_code == 200, response.text
    assert response.headers['content-disposition'].endswith('"vertige-three-statement.xlsx"')
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ['Income statement', 'Balance sheet', 'Cash flow', 'Assumptions']
    income, balance, cash, assumptions = (book[name] for name in book.sheetnames)
    for sheet in book:
        assert sheet.freeze_panes == 'C7' and sheet.sheet_view.showGridLines is False
        assert [sheet.cell(6, c).value for c in range(3, years+3)] == [f'Year {i+1}' for i in range(years)]
        assert sheet.print_title_cols == '$B:$B'
        assert len(sheet.col_breaks.brk) == (years-1)//5
        assert not any(c.data_type in ('f', 'e') for cells in sheet for c in cells)
    def value(sheet, label, col):
        cell = sheet.cell(row(sheet, label), col)
        assert cell.data_type == 'n' and isinstance(cell.value, (int, float)), (label, cell.value)
        assert cell.number_format != 'General'
        return cell.value
    for col in range(3, years+3):
        ni = value(income, 'Net income', col)
        assert ni == pytest.approx(value(income, 'EBIT', col)-value(income, 'Interest expense', col)-value(income, 'Tax expense', col))
        assert value(balance, 'Total assets', col) == pytest.approx(value(balance, 'Total liabilities', col)+value(balance, 'Total equity', col))
        assert value(cash, 'Net income', col) == ni
        assert value(cash, 'Ending cash', col) == value(balance, 'Cash', col)
        assert value(cash, 'Beginning cash', col)+value(cash, 'Net change in cash', col) == pytest.approx(value(cash, 'Ending cash', col))
        assert sum(value(cash, label, col) for label in ['Operating cash flow', 'Investing cash flow', 'Financing cash flow']) == pytest.approx(value(cash, 'Net change in cash', col))
        assert value(cash, 'Capital expenditure', col) < 0
        assert value(cash, 'Debt repaid', col) < 0
        assert value(cash, 'Debt issued', col) == 0
        for sheet, label in [(income, 'EBIT − interest − tax − net income'), (balance, 'Assets − liabilities − equity'), (cash, 'Ending − beginning − net cash change'), (cash, 'CFO + CFI + CFF − net cash change')]:
            assert abs(value(sheet, label, col)) < .01
    assert value(income, 'Revenue', 3) == 1_100_000
    assert value(cash, 'Capital expenditure', 3) == -33_000.5
    assert value(income, 'Depreciation', 3) == pytest.approx(39_567.52375)
    assert value(assumptions, 'Growth capital expenditure', 3) == .5
    assert value(assumptions, 'Receivables / revenue (ratio fallback)', 3) == .12
    assert value(assumptions, 'Days sales outstanding', 3) == 45


def test_period_label_stays_literal_and_missing_days_are_not_zero():
    payload = inputs(1)
    payload['historical_base']['period'] = '=1+1'
    payload['wc_drivers'] = {}
    response = client.post('/api/export/excel/three-statement', json=payload)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content))['Assumptions']
    label = sheet.cell(row(sheet, 'Opening period label'), 3)
    assert label.value == '=1+1' and label.data_type == 's'
    assert sheet.cell(row(sheet, 'Days sales outstanding'), 3).value == 'Use ratio fallback'


@pytest.mark.parametrize('change', ['missing', 'unbalanced', 'wrong_horizon'])
def test_export_rejects_invalid_inputs_instead_of_producing_plausible_workbook(change):
    payload = inputs()
    if change == 'missing':
        del payload['historical_base']['revenue']
    elif change == 'unbalanced':
        payload['historical_base']['total_assets'] += 1_000
    else:
        payload['projection_years'] = 2
    response = client.post('/api/export/excel/three-statement', json=payload)
    assert response.status_code == 422, response.text
