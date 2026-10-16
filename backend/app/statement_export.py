"""Readable, numeric snapshots of the existing three-statement calculation."""
from datetime import datetime, timezone
import math

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from app.engine.three_statement_model import ThreeStatementModel

NAVY, BLUE, GOLD = '0A192F', '637BE6', 'E5B072'
AMOUNT = '#,##0.00;(#,##0.00);"—"'
CHECK = '[>=0.005]#,##0.00;[<=-0.005](#,##0.00);0.00'
PERCENT = '0.0%;(0.0%);"—"'


def _cell(sheet, row, col, value, fmt=AMOUNT, bold=False):
    cell = sheet.cell(row, col)
    if isinstance(value, str):
        # Keep analyst-entered period labels literal, including strings starting '='.
        cell.value, cell.data_type = value, 's'
    else:
        value = float(value)
        if not math.isfinite(value):
            raise ValueError('Statement values exceed Excel numeric limits')
        cell.value = value
        cell.number_format = fmt
    cell.font = Font(name='Arial', size=10, color=NAVY, bold=bold)
    cell.alignment = Alignment(horizontal='left' if isinstance(value, str) else 'right', vertical='center')
    sheet.row_dimensions[row].height = 21
    return cell


def _sheet(book, name, periods, note):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.sheet_properties.tabColor = NAVY if name != 'Assumptions' else BLUE
    sheet.column_dimensions['A'].width = 3
    sheet.column_dimensions['B'].width = 43
    for col in range(3, len(periods) + 3):
        sheet.column_dimensions[get_column_letter(col)].width = 18
    title = _cell(sheet, 2, 2, name)
    title.font = Font(name='Arial', size=14, bold=True, color=NAVY)
    sheet.row_dimensions[2].height = 29
    for col in range(2, len(periods) + 3):
        sheet.cell(2, col).border = Border(bottom=Side(style='thin', color=GOLD))
    for row, text in [(3, 'Input currency units · no scaling or conversion'), (4, note)]:
        _cell(sheet, row, 2, text).font = Font(name='Arial', size=10, italic=True, color='526176')
    for col, label in enumerate(['Line item', *periods], 2):
        cell = _cell(sheet, 6, col, label, bold=True)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='left' if col == 2 else 'center', vertical='center')
    sheet.row_dimensions[6].height = 26
    sheet.freeze_panes = 'C7'
    sheet.print_title_rows = '1:6'
    sheet.print_title_cols = 'B:B'
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 85
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.oddFooter.left.text = 'Vertige Research · Calculated snapshot'
    sheet.oddFooter.right.text = 'Page &P of &N'
    # Long horizons print in groups of five years, with line items repeated.
    for col in range(7, len(periods) + 2, 5):
        sheet.col_breaks.append(Break(id=col))
    return sheet


def _section(sheet, row, label, years):
    for col in range(2, years + 3):
        cell = _cell(sheet, row, col, label if col == 2 else '', bold=True)
        cell.fill = PatternFill('solid', fgColor='EAF0F7')
    sheet.row_dimensions[row].height = 25


def _line(sheet, row, label, values, *, total=False, fmt=AMOUNT, check=False):
    for col, value in enumerate([label, *values], 2):
        cell = _cell(sheet, row, col, value, fmt, total)
        if total:
            cell.border = Border(top=Side(style='thin', color='BEC9D7'))
        elif col == 2 and not check:
            cell.alignment = Alignment(indent=1, vertical='center')
        if col > 2 and isinstance(value, str):
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
            sheet.row_dimensions[row].height = max(21, 15 * math.ceil(len(value) / 21))
        if '%' in fmt or '×' in fmt:
            cell.font = Font(name='Arial', size=10, color=NAVY, italic=True, bold=total)
        if col > 2 and isinstance(value, (int, float)):
            width = len(f'{value:,.2f}') + (2 if value < 0 else 0) + 2
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width, width)
    if check:
        cells = f'C{row}:{get_column_letter(len(values)+2)}{row}'
        sheet.conditional_formatting.add(cells, CellIsRule(
            operator='notBetween', formula=['-0.009999999', '0.009999999'],
            font=Font(color='9C0006', bold=True), fill=PatternFill('solid', fgColor='FCE8E6')))


def statement_workbook(request):
    """Recompute from validated inputs; no client-supplied result or Excel formulas."""
    model = ThreeStatementModel(**{key: getattr(request, key) for key in type(request).model_fields})
    periods = model.project_financials()
    labels = [p.period.replace('_', ' ') for p in periods]
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = 'Vertige Research'
    book.properties.title = 'Three-statement projection'
    book.properties.description = 'Calculated snapshot. Recalculate in Vertige after editing assumptions.'
    layouts = {
        'Income statement': [
            ('Operations', None), ('Revenue', 'revenue', True), ('Cost of goods sold', 'cogs'),
            ('Gross profit', 'gross_profit', True), ('Operating expenses', 'operating_expenses'),
            ('EBITDA', 'ebitda', True), ('Depreciation', 'depreciation'), ('EBIT', 'ebit', True),
            ('Interest expense', 'interest_expense'), ('Earnings before tax', 'ebt', True),
            ('Tax expense', 'tax_expense'), ('Net income', 'net_income', True)],
        'Balance sheet': [
            ('Assets', None), ('Cash', 'cash'), ('Accounts receivable', 'accounts_receivable'),
            ('Inventory', 'inventory'), ('Current assets', 'current_assets', True),
            ('Gross PP&E', 'ppe_gross'), ('Accumulated depreciation', 'accumulated_depreciation'),
            ('Net PP&E', 'ppe_net', True), ('Other assets', 'other_assets'), ('Total assets', 'total_assets', True),
            ('Liabilities', None), ('Accounts payable', 'accounts_payable'), ('Accrued liabilities', 'accrued_liabilities'),
            ('Short-term debt / revolver', 'short_term_debt'), ('Current liabilities', 'current_liabilities', True),
            ('Long-term debt', 'long_term_debt'), ('Other liabilities', 'other_liabilities'), ('Total liabilities', 'total_liabilities', True),
            ('Equity', None), ('Common stock', 'common_stock'), ('Retained earnings', 'retained_earnings'), ('Total equity', 'total_equity', True)],
        'Cash flow': [
            ('Operating activities', None), ('Net income', 'net_income_cf'), ('Depreciation', 'depreciation_cf'),
            ('Receivables: cash effect', 'change_in_ar'), ('Inventory: cash effect', 'change_in_inventory'),
            ('Payables: cash effect', 'change_in_ap'), ('Accruals: cash effect', 'change_in_accrued'), ('Operating cash flow', 'operating_cash_flow', True),
            ('Investing activities', None), ('Capital expenditure', 'capex'), ('Investing cash flow', 'investing_cash_flow', True),
            ('Financing activities', None), ('Debt issued', 'debt_issuance'), ('Debt repaid', 'debt_repayment'),
            ('Dividends paid', 'dividends'), ('Financing cash flow', 'financing_cash_flow', True),
            ('Cash reconciliation', None), ('Beginning cash', 'beginning_cash'), ('Net change in cash', 'net_cash_change', True), ('Ending cash', 'cash', True),
            ('Valuation cash flow', None), ('Unlevered free cash flow', 'unlevered_fcf', True)],
    }
    checks = {
        'Income statement': [('EBIT − interest − tax − net income', lambda p: p.ebit-p.interest_expense-p.tax_expense-p.net_income)],
        'Balance sheet': [('Assets − liabilities − equity', lambda p: p.total_assets-p.total_liabilities-p.total_equity)],
        'Cash flow': [
            ('Ending − beginning − net cash change', lambda p: p.cash-p.beginning_cash-p.net_cash_change),
            ('CFO + CFI + CFF − net cash change', lambda p: p.operating_cash_flow+p.investing_cash_flow+p.financing_cash_flow-p.net_cash_change)],
    }
    for name, layout in layouts.items():
        note = 'Calculated snapshot · relative projection years'
        if name == 'Cash flow':
            note = 'Calculated snapshot · cash outflows in parentheses'
        sheet = _sheet(book, name, labels, note)
        row = 7
        for label, field, *total in layout:
            if field is None:
                _section(sheet, row, label, len(periods))
            else:
                _line(sheet, row, label, [float(getattr(p, field)) for p in periods], total=bool(total))
            row += 1
        _section(sheet, row + 1, 'Reconciliation differences · tolerance < 0.01', len(periods))
        for offset, (label, calculate) in enumerate(checks[name], row + 2):
            _line(sheet, offset, label, [float(calculate(p)) for p in periods], fmt=CHECK, check=True)
        if name == 'Cash flow':
            row = sheet.max_row + 2
            _section(sheet, row, 'Interest coverage covenant', len(periods))
            _line(sheet, row+1, 'EBIT / interest expense', [float(p.ebit/p.interest_expense) if p.interest_expense > 0 else 'No interest expense' for p in periods], fmt='0.00"×"')
            _line(sheet, row+2, 'Minimum coverage', [request.financing_assumptions.covenant_min_interest_coverage]*len(periods), fmt='0.00"×"')
            _line(sheet, row+3, 'Covenant result', ['Below minimum' if p.interest_expense > 0 and p.ebit/p.interest_expense < request.financing_assumptions.covenant_min_interest_coverage else 'No breach' for p in periods])

    sheet = _sheet(book, 'Assumptions', labels, 'Recorded inputs · recalculate in Vertige after edits')
    _section(sheet, 7, 'Annual assumptions', len(periods))
    for row, (label, values, fmt) in enumerate([
        ('Revenue growth', request.operating_assumptions.revenue_growth_rates, PERCENT),
        ('COGS / revenue', request.operating_assumptions.cogs_percent_revenue, PERCENT),
        ('Growth capital expenditure', request.capex_assumptions.growth_capex or [0]*len(periods), AMOUNT)], 8):
        _line(sheet, row, label, values, fmt=fmt)
    row = 12
    groups = [
        ('Operating assumptions', request.operating_assumptions, [
            ('Fixed operating expenses','opex_fixed',AMOUNT),('Variable opex / revenue','opex_variable_pct',PERCENT),('Tax rate','tax_rate',PERCENT)]),
        ('Working capital', request.wc_drivers, [
            ('Receivables / revenue (ratio fallback)','ar_pct_revenue',PERCENT),('Inventory / COGS (ratio fallback)','inventory_pct_cogs',PERCENT),('Payables / COGS (ratio fallback)','ap_pct_cogs',PERCENT),
            ('Days sales outstanding','days_sales_outstanding','0.0'),('Days inventory outstanding','days_inventory_outstanding','0.0'),('Days payable outstanding','days_payable_outstanding','0.0'),
            ('Accrued liabilities / revenue','accrued_liabilities_pct_revenue',PERCENT),('Minimum cash balance','minimum_cash_balance',AMOUNT)]),
        ('Capital expenditure and depreciation', request.capex_assumptions, [
            ('Maintenance capex / revenue','maintenance_capex_pct',PERCENT),('Asset life (years)','asset_life_years','0'),('Salvage value / cost','salvage_value_pct',PERCENT),('Depreciation method','depreciation_method',AMOUNT)]),
        ('Financing', request.financing_assumptions, [
            ('Debt interest rate','debt_interest_rate',PERCENT),('Annual debt amortization','amortization_pct',PERCENT),('Target cash balance','target_cash_balance',AMOUNT),
            ('Dividend payout ratio','dividend_payout_ratio',PERCENT),('Revolver capacity','revolver_capacity',AMOUNT),('Minimum interest coverage','covenant_min_interest_coverage','0.00"×"')]),
    ]
    for heading, source, fields in groups:
        _section(sheet, row, heading, len(periods))
        _cell(sheet, row, 3, 'All years', bold=True)
        row += 1
        for label, key, fmt in fields:
            value = getattr(source, key)
            if value is None:
                value = 'Use ratio fallback'
            elif key == 'depreciation_method':
                value = value.replace('_', ' ').capitalize()
            _line(sheet, row, label, [value], fmt=fmt)
            row += 1
        row += 1
    _section(sheet, row, 'Opening historical inputs', len(periods))
    _cell(sheet, row, 3, 'Opening period', bold=True)
    for key, value in request.historical_base.model_dump().items():
        row += 1
        label = {'period': 'Opening period label', 'ppe_net': 'Net PP&E', 'cogs': 'Cost of goods sold', 'capex': 'Capital expenditure (input sign)'}.get(key, key.replace('_', ' ').capitalize())
        _line(sheet, row, label, [value])
    row += 2
    _section(sheet, row, 'Calculation basis', len(periods))
    for note in [
        'Source: model inputs submitted to Vertige at export.',
        'Opening data may be illustrative; confirm its source.',
        'Days, when entered, override working-capital ratios.',
        'Interest uses average opening and closing debt.',
        'New capex uses half-year depreciation.',
        'Opening net PP&E uses the entered remaining life.',
        'Residual opening assets and liabilities stay constant.',
        'No tax-loss carryforwards are modeled.',
        'Values are snapshots, not linked Excel calculations.',
        'Excel stores approximately 15 significant digits.',
        f'Exported {datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}.',
    ]:
        row += 1
        _cell(sheet, row, 2, note)
    for sheet in book:
        sheet.print_area = f'B2:{get_column_letter(len(periods)+2)}{sheet.max_row}'
    return book
